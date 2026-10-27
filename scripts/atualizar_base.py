"""Atualiza a base do Monitoramento de Fechamento a partir da planilha de controle.

Copia "Resumo da utilização da Conciliação 2.xlsx" (pasta python, acima deste
projeto) para data/monitoramento_fechamento/resumo.xlsx, que é o arquivo que o portal lê,
e grava data/monitoramento_fechamento/status.json com a data da atualização.

Uso:
    python scripts/atualizar_base.py             # só atualiza os arquivos
    python scripts/atualizar_base.py --publicar  # atualiza e faz commit + push

Funciona com a planilha aberta no Excel (lê por uma cópia feita pelo Windows).
"""
import json
import os
import subprocess
import sys
import tempfile
from datetime import datetime
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
PLANILHA = RAIZ.parent / "Resumo da utilização da Conciliação 2.xlsx"
PASTA_DADOS = RAIZ / "data" / "monitoramento_fechamento"
RESUMO = PASTA_DADOS / "resumo.xlsx"
STATUS = PASTA_DADOS / "status.json"

# As etapas ficam fixas no portal (static/script.js); da planilha só vêm os clientes.
ABAS_OBRIGATORIAS = ("Clientes",)


def copiar(origem: Path, destino: Path):
    # O Python não abre um arquivo travado pelo Excel, mas o Copy-Item do Windows copia.
    subprocess.run(
        ["powershell", "-NoProfile", "-Command",
         "Copy-Item -LiteralPath $env:ORIGEM -Destination $env:DESTINO -Force"],
        env={**os.environ, "ORIGEM": str(origem), "DESTINO": str(destino)},
        check=True,
    )


def validar(arquivo: Path) -> int:
    wb = openpyxl.load_workbook(arquivo, read_only=True, data_only=True)
    try:
        faltando = [a for a in ABAS_OBRIGATORIAS if a not in wb.sheetnames]
        if faltando:
            raise SystemExit(f"A planilha não tem as abas: {', '.join(faltando)}")
        linhas = wb["Clientes"].iter_rows(min_row=2, values_only=True)
        return sum(1 for l in linhas if l and l[0] not in (None, ""))
    finally:
        wb.close()


def publicar():
    def git(*args):
        subprocess.run(["git", "-C", str(RAIZ), *args], check=True)

    git("add", str(RESUMO), str(STATUS))
    if subprocess.run(["git", "-C", str(RAIZ), "diff", "--cached", "--quiet"]).returncode == 0:
        print("Nada mudou desde a última publicação.")
        return
    git("commit", "-m", f"Atualização da base: {datetime.now():%Y-%m-%d %H:%M}")
    git("push")


def main():
    if not PLANILHA.exists():
        raise SystemExit(f"Planilha não encontrada: {PLANILHA}")
    PASTA_DADOS.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory() as pasta:
        copia = Path(pasta) / "resumo.xlsx"
        copiar(PLANILHA, copia)
        registros = validar(copia)
        copiar(copia, RESUMO)

    STATUS.write_text(
        json.dumps({"ultima_execucao": datetime.now().isoformat(timespec="seconds"),
                    "registros": registros}, indent=2),
        encoding="utf-8",
    )
    print(f"{registros} clientes gravados em {RESUMO.relative_to(RAIZ)}")

    if "--publicar" in sys.argv:
        publicar()


if __name__ == "__main__":
    main()
